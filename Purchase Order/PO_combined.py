#!/usr/bin/env python3
"""
TCS iON Purchase Order bot — single combined file.

Modes:
  (default)  Create PO from PR (header + line items from Excel/CSV)
  --inspect  Interactive inspector: login, pause for manual navigation,
             then dump DOM / fields / screenshots
  --debug    Automated navigate to PO entry form and dump frames/fields

Usage examples:
  python PO_combined.py
  python PO_combined.py --source data/PO_Input.xlsx
  python PO_combined.py --pr-from 01/08/2026 --pr-to 31/08/2026
  python PO_combined.py --inspect
  python PO_combined.py --debug
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path

# ---------------------------------------------------------------------------
# Path bootstrap (same pattern as original scripts)
# ---------------------------------------------------------------------------
SHARED_DIR = Path(__file__).resolve().parent.parent
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SHARED_DIR.parent
DATA_DIR = PROJECT_DIR / "data"
for _p in (SCRIPT_DIR, SHARED_DIR):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

import requisition as R
from playwright.sync_api import sync_playwright

# Fill helpers shared with the PR filler.
import fill_requisition as F


# ===========================================================================
# PO field maps (from original PO.py)
# ===========================================================================

PO_FIELDS = {
    "po_type": {
        "value_key": "po_type",
        "default": "Domestic",
        "selectors": [
            "select[name='trhtransubtypelviid']",
            "select[name='S_trhpotypeid_trhpotypedcnid']",
            "select[name='trhpotypeid']",
            "select[name='potypeid']",
            "#trhpoTypeid",
            "select:has(option:has-text('Domestic'))",
        ],
    },
    "currency": {
        "value_key": "currency",
        "default": "INR",
        "selectors": [
            "select[name='trhcrncode']",
            "select[name='S_trhcurrencyid_trhcurrencydcnid']",
            "select[name='trhcurrencyid']",
            "select[name='currencyid']",
            "select:has(option:has-text('INR'))",
            "select:has(option:has-text('Rupee'))",
        ],
    },
    "category": {
        "value_key": "category",
        "selectors": [
            "select[name='trhtrancatid_trhtrancatdcnid']",
            "select[name='S_trhtrancatid_trhtrancatdcnid']",
        ],
    },
    "site": {
        "value_key": "site",
        "selectors": [
            "select[name='trhsiteid_trhsitedcnid']",
            "select[name='S_trhsiteid_trhsitedcnid']",
        ],
    },
    "account_site": {
        "value_key": "account_site",
        "selectors": [
            "select[name='trhaccsiteid_trhaccsitedcnid']",
            "select[name='S_trhaccsiteid_trhaccsitedcnid']",
        ],
    },
    "delivery_terms": {
        "value_key": "delivery_terms",
        "selectors": ["select[name='trhdlvtid_trhdlvtdcnid']"],
    },
    "delivery_mode": {
        "value_key": "delivery_mode",
        "selectors": ["select[name='trhdlvmid_trhdlvmdcnid']"],
    },
    "payment_terms": {
        "value_key": "payment_terms",
        "selectors": [
            "select[name='trhpaymentterms']",
            "select[name='S_trhpaymentterms_trhpaymenttermsdcnid']",
            "#trhpaymentterms",
        ],
    },
    "delivery_date": {
        "value_key": "delivery_date",
        "selectors": [
            "input[name='trhdeliverydate']",
            "input[name='trhdate']",
        ],
    },
    "vendor_code": {
        "value_key": "vendor_code",
        "selectors": [
            "input[name='trhpartycode']",
            "[name='trhvendrcode']",
            "[name='partycode']",
        ],
    },
    "vendor_desc": {
        "value_key": "vendor_desc",
        "selectors": [
            "input[name='trhpartydesc']",
            "[name='trhvendrdesc']",
            "[name='partydesc']",
        ],
    },
    "remarks": {
        "value_key": "remarks",
        "selectors": [
            "[name='trhremarks']",
            "[name='trhdesc']",
            "[name='trhcomments']",
        ],
    },
}

PO_ITEM_FIELDS = {
    "item_code": ["#trditemcode{r}", "[name='trditemcode{r}']"],
    "item_desc": ["#trditemdesc{r}", "[name='trditemdesc{r}']"],
    "ac_code": ["#trdacccode{r}", "[name='trdacccode{r}']"],
    "req_date": ["#trdreqdate{r}", "[name='trdreqdate{r}']"],
    "uom": ["#trduomcode{r}", "[name='trduomcode{r}']"],
    "pack_size": ["#trdpacksize{r}", "[name='trdpacksize{r}']"],
    "pack_qty": ["#trdpackquantity{r}", "[name='trdpackquantity{r}']"],
    "qty": ["#trdqty{r}", "[name='trdqty{r}']"],
    "rate": ["#trditemrate{r}", "[name='trditemrate{r}']"],
    "amount": ["#trdamount{r}", "[name='trdamount{r}']"],
    "base_uom": ["#trdbaseuomcode{r}", "[name='trdbaseuomcode{r}']"],
    "base_qty": ["#trdbaseqty{r}", "[name='trdbaseqty{r}']"],
}

PR_COLUMNS = {
    "description": 1,
    "transaction_date": 2,
    "sub_type": 3,
    "category": 4,
    "site": 5,
    "account_site": 6,
    "party_code": 7,
    "party_desc": 8,
    "party_address": 9,
    "remarks": 10,
    "comments": 11,
}

ITEM_COLUMNS = {
    "pr_row": 1,
    "item_code": 2,
    "item_desc": 3,
    "ac_code": 4,
    "req_date": 5,
    "uom": 6,
    "pack_size": 7,
    "pack_qty": 8,
    "qty": 9,
    "base_uom": 10,
    "base_qty": 11,
    "rate": 12,
    "amount": 13,
    "memo": 14,
    "weight_uom": 15,
    "weight_qty": 16,
    "po_number": 17,
    "cost_center": 18,
    "remarks": 19,
    "loc_type": 20,
    "loc_id": 21,
}


# ===========================================================================
# Data loading
# ===========================================================================

def _as_str(value):
    if value is None:
        return ""
    return str(value).strip()


def _detect_flat_headers(headers):
    mapping = {}
    for idx, h in enumerate(headers or []):
        hl = _as_str(h).lower()
        if not hl:
            continue
        if "item" in hl and "code" in hl:
            mapping.setdefault("item_code", idx)
        elif "item" in hl and any(k in hl for k in ("desc", "name", "desic")):
            mapping.setdefault("item_desc", idx)
        elif hl == "qty" or "qty" in hl:
            mapping.setdefault("qty", idx)
        elif "quant" in hl:
            mapping.setdefault("qty", idx)
        elif "delivery" in hl and "date" in hl:
            mapping.setdefault("delivery_date", idx)
        elif "req" in hl and "date" in hl:
            mapping.setdefault("req_date", idx)
        elif "date" in hl:
            mapping.setdefault("req_date", idx)
        elif "po" in hl and "type" in hl:
            mapping.setdefault("po_type", idx)
        elif "currenc" in hl:
            mapping.setdefault("currency", idx)
        elif "payment" in hl or "term" in hl:
            mapping.setdefault("payment_terms", idx)
        elif ("vendor" in hl or "supplier" in hl or "party" in hl) and "code" in hl:
            mapping.setdefault("vendor_code", idx)
        elif ("vendor" in hl or "supplier" in hl or "party" in hl) and any(
            k in hl for k in ("desc", "name")
        ):
            mapping.setdefault("vendor_desc", idx)
        elif "rate" in hl or "price" in hl:
            mapping.setdefault("rate", idx)
        elif "uom" in hl or hl == "unit":
            mapping.setdefault("uom", idx)
    return mapping


def _flat_rows_to_po(rows, headers):
    mapping = _detect_flat_headers(headers)
    if "item_code" not in mapping and "item_desc" not in mapping:
        return None

    def _cell(row, key):
        col = mapping.get(key)
        if col is None or col >= len(row):
            return ""
        return _as_str(row[col])

    header = {}
    items = []
    for idx, row in enumerate(rows, start=1):
        item_code = _cell(row, "item_code")
        item_desc = _cell(row, "item_desc")
        if not item_code and not item_desc:
            continue
        if not item_code and item_desc:
            item_code = re.sub(r"[^A-Z0-9]+", "_", item_desc.upper()).strip("_")[:40] or "ITEM"
        if not item_desc and item_code:
            item_desc = item_code
        uom = _cell(row, "uom") or header.get("uom") or os.getenv("FORM_DEFAULT_UOM", "NOS")
        items.append({
            "item_code": item_code,
            "item_desc": item_desc,
            "ac_code": _cell(row, "ac_code"),
            "req_date": _cell(row, "req_date") or _cell(row, "delivery_date"),
            "uom": uom,
            "pack_size": "",
            "pack_qty": "",
            "qty": _cell(row, "qty") or "1",
            "base_uom": uom,
            "base_qty": _cell(row, "qty") or "1",
            "rate": _cell(row, "rate"),
            "amount": "",
        })
        for key in ("po_type", "currency", "payment_terms", "delivery_date",
                    "vendor_code", "vendor_desc"):
            v = _cell(row, key)
            if v and not header.get(key):
                header[key] = v
    if not items:
        return None
    return {"header": header, "items": items}


def load_po_source(source="PO_Input.xlsx", row=None):
    """Load PO data from Excel or CSV. Returns {"header": {...}, "items": [...]}."""
    path = Path(source)
    if not path.is_absolute():
        stem_parts = list(Path(source).parts)
        if stem_parts and stem_parts[0].lower() == "data":
            path = PROJECT_DIR / path
        else:
            path = DATA_DIR / path
    if not path.is_file():
        if DATA_DIR.is_dir() and path.parent == DATA_DIR:
            wanted = path.name.lower()
            for cand in DATA_DIR.iterdir():
                if cand.name.lower() == wanted:
                    path = cand
                    break
    if not path.is_file():
        print(f"WARNING: source file {path} does not exist.")
        return {}

    suffix = path.suffix.lower()
    if suffix == ".csv":
        return _load_po_csv(path, row)

    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)

    # Legacy two-tab: PR + PR Items / PR Input
    item_sheet = "PR Items" if "PR Items" in wb.sheetnames else (
        "PR Input" if "PR Input" in wb.sheetnames else None
    )
    if "PR" in wb.sheetnames and item_sheet is not None:
        wsi = wb[item_sheet]
        item_rows = []
        for r in wsi.iter_rows(min_row=2, values_only=True):
            if r[0] is None:
                continue
            try:
                pr_idx = int(r[0])
            except Exception:
                continue
            item_rows.append((
                pr_idx,
                {k: (r[c - 1] if (c - 1) < len(r) else "") for k, c in ITEM_COLUMNS.items()},
            ))
        ws = wb["PR"]
        prs = []
        for idx, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if r[0] is None and (len(r) <= 2 or r[2] is None):
                continue
            data = {k: (r[c - 1] if (c - 1) < len(r) else "") for k, c in PR_COLUMNS.items()}
            data["row"] = idx
            data["items"] = [it for i, it in item_rows if i == idx]
            prs.append(data)
        active = [p for p in prs if p.get("items")]
        if active:
            pick = (
                [p for p in active if p["row"] == row]
                if row
                else ([active[row - 1]] if (row and row <= len(active)) else active)
            )
            target = pick[0]
            header = {
                "vendor_code": target.get("party_code") or target.get("vendor_code") or "",
                "vendor_desc": target.get("party_desc") or target.get("vendor_desc") or "",
                "remarks": target.get("remarks") or "",
                "delivery_date": target.get("transaction_date") or "",
            }
            return {"header": header, "items": target["items"]}
        print(f"WARNING: {path.name} PR sheets exist but no rows have items")

    # Flat single-sheet
    for name in wb.sheetnames:
        if name.strip().lower() in {"fields reference", "fields", "reference"}:
            continue
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows or len(rows) < 2:
            continue
        flat = _flat_rows_to_po(rows[1:], rows[0])
        if flat:
            print(f"Loaded PO data from sheet '{name}': {len(flat['items'])} item(s)")
            return flat

    print(f"WARNING: {path.name} has no recognizable PO format")
    return {}


def _load_po_csv(path, row=None):
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        reader = csv.reader(fh)
        all_rows = [r for r in reader if any(cell.strip() for cell in r)]
    if not all_rows or len(all_rows) < 2:
        print(f"WARNING: {path.name} has no data rows")
        return {}
    flat = _flat_rows_to_po(all_rows[1:], all_rows[0])
    if not flat:
        print(f"WARNING: {path.name} has no recognizable PO columns")
        return {}
    print(f"Loaded PO data from {path.name}: {len(flat['items'])} item(s)")
    return flat


# ===========================================================================
# Header / grid filling
# ===========================================================================

def _pick(locators, page):
    for selector in locators:
        loc = F._wait_for(page, selector, max_wait=3)
        if loc is not None:
            try:
                if loc.is_visible():
                    return loc
            except Exception:
                return loc
    return None


def fill_po_header(page, header, status_label="Domestic"):
    header = header or {}
    for key, spec in PO_FIELDS.items():
        value = (header.get(key) or "").strip() or spec.get("default")
        if not value:
            continue
        loc = _pick(spec["selectors"], page)
        if loc is None:
            print(f"  SKIP PO header '{key}': no matching field found")
            continue
        tag = ""
        try:
            tag = loc.evaluate("el => el.tagName.toLowerCase()")
        except Exception:
            pass
        if tag == "select":
            F._select_option(page, spec["selectors"][0] if spec["selectors"] else "", value)
            loc = _pick(spec["selectors"], page)
            if loc is not None:
                try:
                    loc.select_option(label=value)
                except Exception:
                    pass
        else:
            F._set_value(loc, value)
        print(f"  PO header '{key}' = {value!r}")
    return True


def _scroll_to_items(page):
    for selector in ("#insertrow", "#detailBody", "#insertItemRowPR", "button:has-text('Insert Row')"):
        loc = F._wait_for(page, selector, max_wait=2)
        if loc is None:
            continue
        try:
            loc.scroll_into_view_if_needed()
        except Exception:
            pass
        time.sleep(0.3)
        break


def fill_po_items(page, items):
    items = items or []
    if not items:
        print("  No PO items to fill")
        return True
    n = len(items)

    _scroll_to_items(page)

    rows_control = F._wait_for(page, "#insertItemRowPR", max_wait=2)
    if rows_control is not None:
        try:
            rows_control.fill(str(n))
        except Exception:
            pass

    btn = F._wait_for(page, "#insertrow", max_wait=3)
    if btn is None:
        btn = F._wait_for(page, "button:has-text('Insert Row')", max_wait=2)
    if btn is None:
        print("  WARNING: Insert Row button not found; PO grid may be read-only")
        return False
    try:
        btn.scroll_into_view_if_needed()
    except Exception:
        pass
    try:
        btn.click()
    except Exception:
        try:
            btn.click(force=True)
        except Exception:
            pass
    try:
        page.wait_for_load_state("networkidle")
    except Exception:
        pass
    time.sleep(2)

    ok = True
    for i, it in enumerate(items):
        suf = str(i)
        print(f"  Filling PO grid row {i}...")
        F._click_edit_icon(page, i)
        code_loc = _pick([f"#trditemcode{suf}", f"[name='trditemcode{suf}']"], page)
        current_code = ""
        if code_loc is not None:
            try:
                current_code = (code_loc.input_value() or "").strip()
            except Exception:
                current_code = ""
        desired_code = (it.get("item_code") or "").strip()
        if desired_code and current_code and desired_code != current_code:
            print(f"    NOTE row {i}: keeping PR item {current_code!r} (source item {desired_code!r} differs)")
            desired_code = current_code
        elif desired_code and not current_code:
            F._set_value(code_loc, desired_code)
        if desired_code:
            F._click_itembox_sidebutton(page, i)
        for key in ("item_desc", "ac_code", "req_date", "pack_size", "pack_qty",
                    "qty", "rate", "amount", "base_uom", "base_qty"):
            if key not in PO_ITEM_FIELDS:
                continue
            value = (it.get(key) or "").strip()
            if not value:
                continue
            loc = _pick([s.replace("{r}", suf) for s in PO_ITEM_FIELDS[key]], page)
            if loc is None:
                print(f"    SKIP grid field '{key}' row {i}: no matching field")
                continue
            tag = ""
            try:
                tag = loc.evaluate("el => el.tagName.toLowerCase()")
            except Exception:
                pass
            if tag == "select":
                try:
                    loc.select_option(label=value)
                except Exception:
                    pass
            else:
                F._set_value(loc, value)
    return ok


# ===========================================================================
# Create PO from PR (core flow)
# ===========================================================================

def create_po_from_pr(page, data=None, pr_from=None, pr_to=None):
    """Navigate to Purchase Order → Create from PR, fill, and submit."""
    data = data or {}
    header = data.get("header") or {}
    items = data.get("items") or []
    expected_item = (items[0].get("item_code") or "").strip() if items else ""
    _now = datetime.now()
    if not pr_from:
        pr_from = _now.replace(day=1).strftime("%d/%m/%Y")
    if not pr_to:
        pr_to = _now.strftime("%d/%m/%Y")

    # Navigate to Purchase Order
    R._wait_for_item(page, r"Purchase Order")
    po = R._wait_for_item(page, r"Purchase Order", max_wait=10)
    if po is None:
        try:
            page.click("text=Purchase Order")
        except Exception:
            print("Could not find Purchase Order in sidebar")
            return False
    else:
        po.click()
    try:
        page.wait_for_load_state("networkidle", timeout=30000)
    except Exception:
        print("WARNING: networkidle timeout after selecting Purchase Order; continuing")
    page.wait_for_timeout(3000)

    # + Create
    create_btn = None
    for _ in range(30):
        create_btn = R._find_by_role(page, "button", re.compile(r"Create", re.IGNORECASE))
        if create_btn is None:
            create_btn = R._find_sidebar_item(page, r"\+?\s*\bCreate\b")
        if create_btn is not None:
            break
        try:
            page.wait_for_timeout(1000)
        except Exception:
            break
    if create_btn is None:
        print("FAIL: could not find + Create button")
        return False
    if not R._click_robust(create_btn):
        print("FAIL: could not click + Create button")
        return False
    page.wait_for_timeout(1500)

    # Create from PR
    from_pr = R._wait_for_visible(
        page, re.compile(r"Create from PR", re.IGNORECASE),
        tags="li,a,button,span", max_wait=10,
    )
    if from_pr is None:
        print("FAIL: could not find 'Create from PR' option")
        try:
            page.screenshot(path=f"{R.SCREENSHOT_DIR}/po_create_from_pr_not_found.png", full_page=True)
        except Exception:
            pass
        return False
    if not R._click_robust(from_pr):
        print("FAIL: could not click 'Create from PR'")
        return False
    try:
        page.wait_for_load_state("networkidle")
    except Exception:
        pass
    page.wait_for_timeout(1500)

    # Status = Approved
    status_set = False
    for sel in ("select[name='trhstatuslviid']", "select[name*='status']"):
        loc = F._wait_for(page, sel, max_wait=5)
        if loc is not None:
            try:
                loc.select_option(label="Approved")
                status_set = True
                print("  Status set to Approved")
                break
            except Exception:
                try:
                    loc.select_option(index=0)
                    status_set = True
                    print("  Status set via index 0")
                    break
                except Exception:
                    pass
    if not status_set:
        print("WARNING: Could not set Status filter via known selectors")

    # Date range
    for name, val in [("s_FromDate", pr_from), ("s_ToDate", pr_to)]:
        loc = F._wait_for(page, f"input[name='{name}']", max_wait=3)
        if loc is not None:
            try:
                loc.fill(val)
                print(f"  Set {name} = {val}")
            except Exception:
                F._set_value(loc, val)

    # Item code filter
    if expected_item:
        item_loc = F._wait_for(page, "input[name='itemcode']", max_wait=3)
        if item_loc is None:
            item_loc = F._wait_for(page, "#satitemcode", max_wait=3)
        if item_loc is not None:
            F._set_value(item_loc, expected_item)
            print(f"  Set Item Code filter = {expected_item}")
        else:
            print("  WARNING: Item Code filter field not found; search not narrowed")

    # Search
    search_btn = F._wait_for(page, "button:has-text('Search')", max_wait=5)
    if search_btn is None:
        search_btn = F._wait_for(page, "button:text-is('Search')", max_wait=3)
    if search_btn is None:
        for scope in F._scopes(page):
            try:
                btns = scope.locator("button").all()
                for b in btns:
                    try:
                        if b.is_visible() and "search" in (b.inner_text() or "").lower():
                            search_btn = b
                            break
                    except Exception:
                        continue
                if search_btn:
                    break
            except Exception:
                continue
    if search_btn is None:
        print("FAIL: Search button not found")
        return False
    try:
        search_btn.click()
    except Exception:
        search_btn.click(force=True)
    try:
        page.wait_for_load_state("networkidle")
    except Exception:
        pass
    time.sleep(3)

    # Results table
    table_found = False
    for sel in ("table tbody tr", "table tr", "table"):
        loc = F._wait_for(page, sel, max_wait=10)
        if loc is not None:
            try:
                if loc.count() > 0:
                    table_found = True
                    break
            except Exception:
                pass
    if not table_found:
        print("WARNING: No results table found after search")

    # Select matching PR row
    checkbox = None
    if expected_item:
        for scope in F._scopes(page):
            try:
                rows = scope.locator("table tbody tr")
                for i in range(rows.count()):
                    row = rows.nth(i)
                    try:
                        if not row.is_visible():
                            continue
                        txt = (row.inner_text(timeout=1500) or "")
                        if expected_item in txt and row.locator("input[type='checkbox']").count() > 0:
                            checkbox = row.locator("input[type='checkbox']").first
                            print(f"  Found PR row for item {expected_item}: {txt[:120]!r}")
                            break
                    except Exception:
                        continue
                if checkbox is not None:
                    break
            except Exception:
                continue
        if checkbox is None:
            print(
                f"FAIL: no approved PR found for item {expected_item} in {pr_from}..{pr_to}. "
                "PO not created (avoids PO for a different item)."
            )
            for scope in F._scopes(page):
                try:
                    rows = scope.locator("table tbody tr")
                    for i in range(min(rows.count(), 20)):
                        try:
                            txt = (rows.nth(i).inner_text(timeout=1500) or "").replace("\n", " | ")[:250]
                            if txt.strip():
                                print(f"  RESULT ROW {i}: {txt}")
                        except Exception:
                            continue
                except Exception:
                    continue
            try:
                page.screenshot(path=f"{R.SCREENSHOT_DIR}/po_search_no_match.png", full_page=True)
                print(f"  Screenshot: {R.SCREENSHOT_DIR}/po_search_no_match.png")
            except Exception:
                pass
            return False

    if checkbox is None:
        for sel in (
            "table tbody tr:first-child input[type='checkbox']",
            "table tbody tr:first-child td input[type='checkbox']",
            "table tbody tr:first-child td:first-child input",
        ):
            checkbox = F._wait_for(page, sel, max_wait=3)
            if checkbox is not None:
                break
    if checkbox is None:
        for scope in F._scopes(page):
            try:
                cbs = scope.locator("table input[type='checkbox']")
                if cbs.count() > 0:
                    checkbox = cbs.first
                    break
            except Exception:
                continue
    if checkbox is None:
        print("FAIL: Could not find checkbox to select PR row")
        return False
    try:
        checkbox.click()
    except Exception:
        checkbox.click(force=True)
    print("  Selected PR row")

    # Create PO
    create_po_btn = None
    for sel in (
        "button:has-text('Create PO')",
        "button:text-is('Create PO')",
        "input[value='Create PO']",
    ):
        create_po_btn = F._wait_for(page, sel, max_wait=3)
        if create_po_btn is not None:
            break
    if create_po_btn is None:
        for scope in F._scopes(page):
            try:
                btns = scope.locator("button, input[type='submit'], input[type='button']").all()
                for b in btns:
                    try:
                        label = (b.inner_text() or b.get_attribute("value") or "").lower()
                        if b.is_visible() and "create po" in label:
                            create_po_btn = b
                            break
                    except Exception:
                        continue
                if create_po_btn:
                    break
            except Exception:
                continue
    if create_po_btn is None:
        print("FAIL: Create PO button not found")
        return False
    try:
        create_po_btn.click()
    except Exception:
        create_po_btn.click(force=True)
    try:
        page.wait_for_load_state("networkidle")
    except Exception:
        pass
    time.sleep(3)
    print("  Clicked Create PO")

    # Fill header + items
    fill_po_header(page, header)
    fill_po_items(page, items)

    # Verify item codes
    mismatch = False
    for i, it in enumerate(items):
        expected_code = (it.get("item_code") or "").strip()
        if not expected_code:
            continue
        loc = _pick([f"#trditemcode{i}", f"[name='trditemcode{i}']"], page)
        if loc is None:
            print(f"  VERIFY FAIL row {i}: item code field not found")
            mismatch = True
            continue
        try:
            actual = (loc.input_value() or "").strip()
        except Exception:
            actual = ""
        if actual != expected_code:
            print(f"  VERIFY FAIL row {i}: item code expected {expected_code!r} got {actual!r}")
            mismatch = True
        else:
            print(f"  VERIFY OK row {i}: item code {actual!r}")
    if mismatch:
        print("FAIL: PO grid item does not match source/PR item; not submitting PO")
        return False

    page.screenshot(path=f"{R.SCREENSHOT_DIR}/po_filled.png", full_page=True)

    # Submit
    submit_btn = F._wait_for(page, "button:has-text('Submit')", max_wait=10)
    if submit_btn is None:
        submit_btn = F._wait_for(page, "button:text-is('Submit')", max_wait=5)
    if submit_btn is None:
        for scope in F._scopes(page):
            try:
                btns = scope.locator("button").all()
                for b in btns:
                    try:
                        if b.is_visible() and "submit" in (b.inner_text() or "").lower():
                            submit_btn = b
                            break
                    except Exception:
                        continue
                if submit_btn:
                    break
            except Exception:
                continue
    if submit_btn is None:
        print("FAIL: Submit button not found")
        return False
    try:
        submit_btn.scroll_into_view_if_needed()
    except Exception:
        pass
    time.sleep(1)
    try:
        submit_btn.click()
    except Exception:
        try:
            submit_btn.click(force=True)
        except Exception:
            try:
                submit_btn.evaluate("el => el.click()")
            except Exception:
                print("FAIL: could not click Submit button")
                return False
    try:
        page.wait_for_load_state("networkidle")
    except Exception:
        pass
    time.sleep(5)

    print("SUCCESS: Purchase Order created successfully from PR")
    return True


# ===========================================================================
# Debug / Inspect helpers (merged from debug_*.py + inspect_po_form.py)
# ===========================================================================

FIELD_JS = """() => {
  const out = [];
  const labels = {};
  document.querySelectorAll('label').forEach(l => {
    if (l.htmlFor) labels[l.htmlFor] = (l.innerText || '').trim();
  });
  const controls = document.querySelectorAll('input:not([type=hidden]), select, textarea');
  for (const el of controls) {
    const rect = el.getBoundingClientRect();
    if (rect.bottom <= 0 || rect.right <= 0) continue;
    const tag = el.tagName.toLowerCase();
    const rec = {
      tag,
      type: el.type || '',
      id: el.id || '',
      name: el.name || '',
      placeholder: el.placeholder || '',
      aria_label: el.getAttribute('aria-label') || '',
      label: labels[el.id] || el.getAttribute('aria-label') || el.placeholder || '',
      value: el.value || '',
      required: !!(el.required || el.getAttribute('aria-required') === 'true'),
      disabled: !!el.disabled,
    };
    if (tag === 'select') {
      rec.selected = el.value || '';
      rec.options = Array.from(el.options).map(o => (o.text || '').trim()).slice(0, 80);
    }
    if (tag === 'textarea') rec.rows = el.rows;
    out.push(rec);
  }
  return out;
}"""

WIDGET_JS = """() => {
  const out = [];
  const seen = new Set();
  const els = document.querySelectorAll(
    '[role="combobox"], [role="listbox"], [aria-haspopup="listbox"], ' +
    '[class*="dropdown-toggle"], [class*="btn-select"], [class*="p-select"]'
  );
  for (const el of els) {
    const rect = el.getBoundingClientRect();
    if (rect.width === 0 || rect.height === 0) continue;
    const text = (el.innerText || '').trim().replace(/\\s+/g, ' ').slice(0, 80);
    const key = el.id + '|' + text;
    if (seen.has(key)) continue;
    seen.add(key);
    out.push({
      role: el.getAttribute('role') || el.tagName.toLowerCase(),
      id: el.id || '',
      text,
      title: el.getAttribute('title') || '',
      aria_expanded: el.getAttribute('aria-expanded') || '',
      class: (el.className || '').toString().slice(0, 80),
    });
  }
  return out;
}"""

TABLE_JS = """() => {
  const out = [];
  const seen = new Set();
  for (const t of document.querySelectorAll('table')) {
    const rect = t.getBoundingClientRect();
    if (rect.width === 0 || rect.height === 0) continue;
    const headers = Array.from(t.querySelectorAll('th'))
      .map(th => (th.innerText || '').trim()).filter(Boolean).slice(0, 40);
    const key = headers.join('|');
    if (seen.has(key)) continue;
    seen.add(key);
    out.push({ headers, rows: t.querySelectorAll('tbody tr').length });
  }
  return out;
}"""

GRID_FIELD_JS = """() => {
  const out = [];
  const labels = {};
  document.querySelectorAll('label').forEach(l => {
    if (l.htmlFor) labels[l.htmlFor] = (l.innerText || '').trim();
  });
  for (const t of document.querySelectorAll('table')) {
    for (const el of t.querySelectorAll('input:not([type=hidden]):not([type=button]), select, textarea')) {
      const rect = el.getBoundingClientRect();
      if (rect.width === 0 && rect.height === 0) continue;
      const tag = el.tagName.toLowerCase();
      out.push({
        tag,
        type: el.type || '',
        id: el.id || '',
        name: el.name || '',
        placeholder: el.placeholder || '',
        aria_label: el.getAttribute('aria-label') || '',
        label: labels[el.id] || el.getAttribute('aria-label') || el.placeholder || '',
        value: el.value || '',
        required: !!(el.required || el.getAttribute('aria-required') === 'true'),
      });
    }
  }
  return out;
}"""


def _collect(scope, js):
    try:
        return scope.evaluate(js) or []
    except Exception as exc:
        print(f"  (frame collect error: {exc})")
        return []


def dump_page_structure(page, label="PO form"):
    """Print frames, selects, inputs, buttons, tables (from debug scripts)."""
    print(f"\n===== {label} =====")
    print(f"URL: {page.url}")
    print(f"Total frames: {len(page.frames)}")
    for i, f in enumerate(page.frames):
        print(f"  [{i}] name={f.name!r} url={f.url[:120]}")
        try:
            print(
                f"       selects={f.locator('select').count()} "
                f"inputs={f.locator('input').count()} "
                f"buttons={f.locator('button').count()} "
                f"tables={f.locator('table').count()}"
            )
        except Exception:
            pass

    for i, f in enumerate(page.frames):
        try:
            sels = f.locator("select").all()
            vis = [s for s in sels if s.is_visible()]
            if vis:
                print(f"\n--- Frame {i} selects ---")
                for s in vis:
                    try:
                        nm = s.get_attribute("name") or s.get_attribute("id") or "?"
                        opts = [o.inner_text()[:30] for o in s.locator("option").all()[:6]]
                        print(f"  name={nm} opts={opts}")
                    except Exception:
                        pass
        except Exception:
            pass

    for i, f in enumerate(page.frames):
        try:
            inps = f.locator("input").all()
            vis = [x for x in inps if x.is_visible()]
            if vis:
                print(f"\n--- Frame {i} visible inputs ---")
                for inp in vis[:25]:
                    try:
                        nm = inp.get_attribute("name") or inp.get_attribute("id") or "?"
                        tp = inp.get_attribute("type") or "text"
                        pl = inp.get_attribute("placeholder") or ""
                        print(f"  name={nm} type={tp} placeholder={pl}")
                    except Exception:
                        pass
        except Exception:
            pass

    for i, f in enumerate(page.frames):
        try:
            btns = f.locator("button").all()
            vis = [x for x in btns if x.is_visible()]
            if vis:
                print(f"\n--- Frame {i} visible buttons ---")
                for btn in vis[:20]:
                    try:
                        t = btn.inner_text()[:50].strip()
                        print(f"  {t!r}")
                    except Exception:
                        pass
        except Exception:
            pass

    for i, f in enumerate(page.frames):
        try:
            tables = f.locator("table").all()
            vis = [x for x in tables if x.is_visible()]
            if vis:
                print(f"\n--- Frame {i} tables ---")
                for tbl in vis[:3]:
                    try:
                        rows = tbl.locator("tr").count()
                        hdr = tbl.locator("tr").first.inner_text()[:300]
                        print(f"  rows={rows} header={hdr!r}")
                    except Exception:
                        pass
        except Exception:
            pass

    print("\n--- PO keywords in page text ---")
    for i, f in enumerate(page.frames):
        try:
            body = f.locator("body").first.inner_text(timeout=3000)
            for kw in [
                "PO Type", "Currency", "Vendor", "Submit", "Save",
                "Payment", "Delivery", "Item Code", "Quantity",
            ]:
                if kw.lower() in body.lower():
                    print(f'  Frame {i}: found "{kw}"')
        except Exception:
            pass


def inspect_po_form(page):
    """Full structured dump → data/po_form_fields.json / .txt / .html."""
    data = {"page_url": page.url, "frames": []}
    dom_parts = []

    for i, scope in enumerate(R._frame_scopes(page)):
        try:
            url = scope.url
        except Exception:
            url = "?"
        fields = _collect(scope, FIELD_JS)
        widgets = _collect(scope, WIDGET_JS)
        tables = _collect(scope, TABLE_JS)
        data["frames"].append({
            "index": i,
            "url": url,
            "fields": fields,
            "widgets": widgets,
            "tables": tables,
        })
        try:
            dom_parts.append(f"<!-- FRAME {i}: {url} -->\n" + scope.content())
        except Exception as exc:
            dom_parts.append(f"<!-- FRAME {i}: content unavailable: {exc} -->\n")

    try:
        page.screenshot(path=f"{R.SCREENSHOT_DIR}/po_inspect_form.png", full_page=True)
    except Exception:
        pass

    # Try Insert Row to expose grid
    grid_fields = []
    for scope in R._frame_scopes(page):
        for btn_sel in [
            "input#insertrow",
            "input[name='button2']",
            "#insertrow",
            "button:has-text('Insert Row')",
            "a:has-text('Insert Row')",
        ]:
            try:
                btn = scope.locator(btn_sel).first
                if btn.count() == 0 or not btn.is_visible():
                    continue
                btn.click()
                try:
                    page.wait_for_load_state("networkidle")
                except Exception:
                    pass
                page.wait_for_timeout(2000)
                grid_fields = _collect(scope, GRID_FIELD_JS)
                if grid_fields:
                    print(f"  Grid fields found via {btn_sel}: {len(grid_fields)} fields")
                break
            except Exception:
                continue
        if grid_fields:
            break
    if grid_fields:
        data["grid_fields"] = grid_fields

    try:
        page.screenshot(path=f"{R.SCREENSHOT_DIR}/po_inspect_final.png", full_page=True)
    except Exception:
        pass

    os.makedirs(DATA_DIR, exist_ok=True)

    json_path = DATA_DIR / "po_form_fields.json"
    with json_path.open("w", encoding="utf-8") as fh:
        json.dump(data, fh, ensure_ascii=False, indent=2)
    print(f"Wrote {json_path.name}")

    dom_path = DATA_DIR / "po_form_dom.html"
    with dom_path.open("w", encoding="utf-8") as fh:
        fh.write("\n".join(dom_parts))
    print(f"Wrote {dom_path.name}")

    lines = ["PURCHASE ORDER FORM INVENTORY", f"URL: {page.url}", ""]
    for fr in data["frames"]:
        lines.append(f"--- Frame {fr['index']}: {fr['url']} ---")
        for t in fr["tables"]:
            lines.append(f"  TABLE (rows={t['rows']}) headers: {t['headers']}")
        for f in fr["fields"]:
            req = " *" if f["required"] else ""
            sel = f" selected={f['selected']!r}" if f.get("selected") else ""
            opts = f" options={f['options']}" if f.get("options") else ""
            lines.append(
                f"  FIELD [{f['tag']}{':' + f['type'] if f['type'] else ''}] "
                f"{f['label'] or f['placeholder'] or f['id'] or f['name']}{req}"
                f"  id={f['id']!r} name={f['name']!r} value={f['value']!r}{sel}{opts}"
            )
        for w in fr["widgets"]:
            lines.append(
                f"  WIDGET [{w['role']}] {w['text']}  id={w['id']!r} "
                f"expanded={w['aria_expanded']!r} class={w['class']!r}"
            )
        lines.append("")
    if data.get("grid_fields"):
        lines.append("--- ITEM GRID FIELDS (after Insert Row) ---")
        for f in data["grid_fields"]:
            lines.append(
                f"  GRID [{f['tag']}{':' + f['type'] if f['type'] else ''}] "
                f"{f['label'] or f['placeholder'] or f['id'] or f['name']}"
                f"  id={f['id']!r} name={f['name']!r} value={f['value']!r}"
            )
        lines.append("")

    txt_path = DATA_DIR / "po_form_fields.txt"
    with txt_path.open("w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))
    print(f"Wrote {txt_path.name} ({len(lines)} lines)")
    print("\n" + "\n".join(lines))

    all_fields = [f for fr in data["frames"] for f in fr["fields"]]
    all_widgets = [w for fr in data["frames"] for w in fr["widgets"]]
    print(
        f"\nSummary: {len(all_fields)} fields, {len(all_widgets)} widgets, "
        f"{len(data.get('grid_fields', []))} grid fields across "
        f"{len(data['frames'])} frames"
    )


def navigate_to_po_entry_form(page, pr_from="01/01/2026", pr_to="31/12/2026"):
    """Automated path used by --debug: go all the way to the PO entry form."""
    proc = R._wait_for_item(page, r"Procurement", max_wait=10)
    if proc:
        proc.click()
    time.sleep(3)
    po = R._wait_for_item(page, r"Purchase Order", max_wait=15)
    if not po:
        print("FAIL: Purchase Order not found")
        return False
    po.click()
    time.sleep(3)
    try:
        page.wait_for_load_state("networkidle", timeout=30000)
    except Exception:
        pass

    cb = R._find_by_role(page, "button", re.compile(r"Create", re.IGNORECASE))
    if cb is None:
        cb = R._find_sidebar_item(page, r"\+?\s*\bCreate\b")
    if not cb:
        print("FAIL: Create button not found")
        return False
    R._click_robust(cb)
    time.sleep(2)

    fp = R._wait_for_visible(
        page, re.compile(r"Create from PR", re.IGNORECASE),
        tags="li,a,button,span", max_wait=10,
    )
    if not fp:
        print("FAIL: Create from PR not found")
        return False
    R._click_robust(fp)
    time.sleep(3)
    try:
        page.wait_for_load_state("networkidle", timeout=30000)
    except Exception:
        pass

    s = F._wait_for(page, "select[name='trhstatuslviid']", max_wait=5)
    if s:
        s.select_option(label="Approved")
    fd = F._wait_for(page, "input[name='s_FromDate']", max_wait=3)
    if fd:
        F._set_value(fd, pr_from)
    td = F._wait_for(page, "input[name='s_ToDate']", max_wait=3)
    if td:
        F._set_value(td, pr_to)
    sb = F._wait_for(page, "button:has-text('Search')", max_wait=5)
    if sb:
        sb.click()
    time.sleep(3)
    try:
        page.wait_for_load_state("networkidle", timeout=30000)
    except Exception:
        pass
    time.sleep(2)

    cb2 = F._wait_for(page, "table tbody tr:first-child input[type='checkbox']", max_wait=5)
    if cb2:
        cb2.click()
    time.sleep(1)
    cp = F._wait_for(page, "button:has-text('Create PO')", max_wait=5)
    if cp:
        cp.click()
    time.sleep(8)
    try:
        page.wait_for_load_state("networkidle", timeout=60000)
    except Exception:
        pass
    time.sleep(3)
    return True


# ===========================================================================
# JSON PO MODE (worker contract: create_po)
# ===========================================================================

PO_NO_PATTERNS = [
    re.compile(r"([A-Z]{2,6}\s*/\s*\d{2,4}\s*/\s*PO\s*/\s*\d{3,})", re.IGNORECASE),
    re.compile(r"([A-Z]{2,6}\s*/\s*\d{4}\s*/\s*\d{3,})", re.IGNORECASE),
    re.compile(r"PO\s*(?:No|Number|#)?\s*[:.\-]?\s*([A-Z0-9][A-Z0-9/\-]{5,})", re.IGNORECASE),
]


def _capture_po_number(page) -> str | None:
    """Best-effort read of the generated PO number from the page after submit."""
    texts = []
    try:
        for scope in R._frame_scopes(page):
            try:
                texts.append(scope.locator("body").first.inner_text(timeout=2000))
            except Exception:
                continue
            try:
                for inp in scope.locator(
                    "input[type='text'], input:not([type]), input[name], textarea"
                ).all():
                    try:
                        v = inp.input_value()
                        if isinstance(v, str) and v.strip():
                            texts.append(v)
                    except Exception:
                        continue
            except Exception:
                continue
    except Exception:
        pass
    if not texts:
        try:
            texts.append(page.inner_text("body", timeout=2000))
        except Exception:
            pass
    blob = "\n".join(texts)
    for pat in PO_NO_PATTERNS:
        for m in pat.finditer(blob):
            val = m.group(1).strip()
            if "/PR/" in val.upper():
                continue
            return val
    return None


def build_po_from_json(payload: dict) -> dict:
    """Map the create_po JSON payload to PO data used by create_po_from_pr."""
    item_code = str(payload.get("itemCode") or payload.get("item_code") or "").strip()
    if not item_code:
        raise ValueError("itemCode is required")
    qty = str(payload.get("qty") or payload.get("quantity") or 1)
    uom = str(payload.get("uom") or "").strip() or os.getenv("FORM_DEFAULT_UOM", "NOS")
    vendor = str(payload.get("vendorCode") or payload.get("vendor_code") or "").strip()
    return {
        "pr_number": str(payload.get("prNumber") or payload.get("pr_number") or "").strip(),
        "header": {"vendor_code": vendor},
        "items": [{
            "item_code": item_code,
            "item_desc": item_code,
            "qty": qty,
            "uom": uom,
            "base_uom": uom,
            "base_qty": qty,
        }],
    }


def _load_json_input(path: str) -> dict:
    p = Path(path)
    if not p.is_absolute():
        for base in (PROJECT_DIR, SCRIPT_DIR, DATA_DIR):
            cand = base / p
            if cand.is_file():
                p = cand
                break
    if not p.is_file():
        raise SystemExit(f"JSON input file not found: {path}")
    return json.loads(p.read_text(encoding="utf-8"))


def run_json_po(payload: dict) -> dict:
    """Run a single PO from the create_po JSON contract. Result JSON = last stdout line."""
    started = time.time()
    pr_number = str(payload.get("prNumber") or payload.get("pr_number") or "").strip()
    vendor = str(payload.get("vendorCode") or payload.get("vendor_code") or "").strip()
    item_code = str(payload.get("itemCode") or payload.get("item_code") or "").strip()
    request_id = str(payload.get("requestId") or payload.get("request_id") or "").strip()
    qty = payload.get("qty") or payload.get("quantity") or 1

    result = {
        "ok": False,
        "action": "create_po",
        "poNumber": None,
        "prNumber": pr_number,
        "vendorCode": vendor,
        "status": None,
        "requestId": request_id,
        "durationMs": None,
        "error": None,
    }

    def _finish(ok, po_no=None, status=None, error=None):
        result["ok"] = ok
        result["poNumber"] = po_no
        result["status"] = status
        result["error"] = str(error)[:500] if error else None
        if status is None and ok:
            result["status"] = "created"
        result["durationMs"] = int((time.time() - started) * 1000)
        return result

    try:
        data = build_po_from_json(payload)
    except Exception as exc:
        return _finish(False, error=str(exc))

    os.makedirs(R.SCREENSHOT_DIR, exist_ok=True)
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False, args=["--start-maximized"])
        context = browser.new_context(no_viewport=True)
        page = context.new_page()
        try:
            if not R.do_login(page):
                return _finish(False, error="login failed")
            mfg_page = R.select_manufacturing(page, enter_requisition=False)
            if mfg_page is None:
                return _finish(False, error="could not reach Manufacturing/Procurement")
            ok = create_po_from_pr(mfg_page, data)
            if not ok:
                return _finish(False, error="PO creation failed in TCS portal")
            po_no = _capture_po_number(mfg_page)
            if po_no:
                return _finish(True, po_no=po_no)
            return _finish(
                False,
                error=(
                    "PO submitted but number could not be captured from the page; "
                    "verify in the PO list before retrying to avoid duplicates"
                ),
            )
        except Exception as exc:
            return _finish(False, error=str(exc))
        finally:
            try:
                R.safe_logout(page)
            except Exception:
                pass
            try:
                browser.close()
            except Exception:
                pass


# ===========================================================================
# Entry points
# ===========================================================================

def run_create(args) -> int:
    default_source = str(DATA_DIR / "PO_Input.xlsx")
    source = args.source or (default_source if Path(default_source).is_file() else None)
    data = {}
    if source:
        data = load_po_source(source, row=args.row)

    os.makedirs(R.SCREENSHOT_DIR, exist_ok=True)
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False, args=["--start-maximized"])
        context = browser.new_context(no_viewport=True)
        page = context.new_page()
        try:
            if not R.do_login(page):
                print("FAIL: login failed")
                return 1
            mfg_page = R.select_manufacturing(page, enter_requisition=False)
            if mfg_page is None:
                print("FAIL: could not reach Manufacturing/Procurement")
                return 1
            if not create_po_from_pr(mfg_page, data, pr_from=args.pr_from, pr_to=args.pr_to):
                print("FAIL: PO creation failed")
                return 1
        except KeyboardInterrupt:
            print("Interrupted by user.")
            return 130
        finally:
            R.safe_logout(page)
            try:
                browser.close()
            except Exception:
                pass
    return 0


def run_debug(args) -> int:
    os.makedirs(R.SCREENSHOT_DIR, exist_ok=True)
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False, args=["--start-maximized"])
        context = browser.new_context(no_viewport=True)
        page = context.new_page()
        try:
            if not R.do_login(page):
                print("FAIL: login")
                return 1
            mfg = R.select_manufacturing(page, enter_requisition=False)
            if not mfg:
                print("FAIL: mfg")
                return 1
            ok = navigate_to_po_entry_form(
                mfg,
                pr_from=args.pr_from or "01/01/2026",
                pr_to=args.pr_to or "31/12/2026",
            )
            if not ok:
                return 1
            dump_page_structure(mfg, label="PO ENTRY FORM")
        finally:
            R.safe_logout(page)
            try:
                browser.close()
            except Exception:
                pass
    return 0


def run_inspect(args) -> int:
    os.makedirs(R.SCREENSHOT_DIR, exist_ok=True)
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False, args=["--start-maximized"])
        context = browser.new_context(no_viewport=True)
        page = context.new_page()
        try:
            if not R.do_login(page):
                print("FAIL: login failed")
                return 1
            try:
                page.screenshot(path=f"{R.SCREENSHOT_DIR}/po_inspect_postlogin.png", full_page=True)
            except Exception:
                pass

            print("\n" + "=" * 70)
            print("  PO FORM INSPECTOR")
            print("=" * 70)
            print()
            print("  The browser is now logged into TCS iON.")
            print("  Please navigate to the Purchase Order Create form manually.")
            print()
            print("  Typical path:")
            print("    Home → Manufacturing → Procurement → Purchase Order → + Create")
            print()
            print("  Once the PO form is fully loaded and visible, press ENTER.")
            print("=" * 70)
            input("\n>>> Press ENTER when the PO form is visible... ")

            try:
                page.wait_for_load_state("networkidle")
            except Exception:
                pass
            page.wait_for_timeout(2000)

            print("\nInspecting PO form...")
            inspect_po_form(page)
            print("\nDone. You can close the browser or press Ctrl+C to exit.")
            try:
                input("\n>>> Press ENTER to close the browser and exit... ")
            except (EOFError, KeyboardInterrupt):
                pass
        except KeyboardInterrupt:
            print("\nInterrupted by user.")
            return 130
        finally:
            R.safe_logout(page)
            try:
                browser.close()
            except Exception:
                pass
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(
        description="TCS iON Purchase Order bot (Create from PR + debug/inspect)"
    )
    parser.add_argument(
        "--source",
        default=None,
        help="Excel/CSV file with PO items. Defaults to data/PO_Input.xlsx if present.",
    )
    parser.add_argument(
        "--row",
        type=int,
        default=None,
        help="For legacy two-tab Excel: which PR row to create the PO from.",
    )
    parser.add_argument(
        "--json",
        default=None,
        help="JSON input file (create_po contract: action/prNumber/vendorCode/itemCode/qty/requestId)",
    )
    parser.add_argument(
        "--pr-from",
        default=None,
        help="PR 'from' date (dd/MM/yyyy). Default: 1st of current month.",
    )
    parser.add_argument(
        "--pr-to",
        default=None,
        help="PR 'to' date (dd/MM/yyyy). Default: today.",
    )
    parser.add_argument(
        "--debug",
        action="store_true",
        help="Automated navigate to PO entry form and dump structure.",
    )
    parser.add_argument(
        "--inspect",
        action="store_true",
        help="Interactive inspector: pause for manual navigation, then dump form.",
    )
    args = parser.parse_args()

    if args.json:
        payload = _load_json_input(args.json)
        result = run_json_po(payload)
        print(json.dumps(result, ensure_ascii=False))
        return 0 if result.get("ok") else 1

    if args.inspect:
        return run_inspect(args)
    if args.debug:
        return run_debug(args)
    return run_create(args)


if __name__ == "__main__":
    sys.exit(main())
