#!/usr/bin/env python3
"""
TCS iON Purchase Requisition (PR) bot — single combined file.

Modes:
  --item-code / --qty     One-shot: create + approve one PR for a given item
  --source FILE|gsheet    Batch pipeline from Excel or Google Sheet
  --request-id TOKEN      Form-driven: load one Requests row and create PR
  --worker                Poll Sheets queue (approved rows with pending TCS status)

Examples:
  python PR_combined.py --item-code PCPWB60132 --qty 50
  python PR_combined.py --source data/PR_Input.xlsx
  python PR_combined.py --source gsheet --row 3
  python PR_combined.py --request-id REQ-001
  python PR_combined.py --worker --once
  python PR_combined.py --worker --interval 60
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
import traceback
from datetime import datetime
from pathlib import Path

# ---------------------------------------------------------------------------
# Paths
# ---------------------------------------------------------------------------
SHARED_DIR = Path(__file__).resolve().parent.parent
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SHARED_DIR.parent
DATA_DIR = PROJECT_DIR / "data"
LOG_DIR = PROJECT_DIR / "logs"
SCREENSHOT_DIR = LOG_DIR / "screenshots"
FORM_LOG = LOG_DIR / "form_pr_log.jsonl"

for _p in (SCRIPT_DIR, SHARED_DIR):
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))

# Load .env early (USERNAME/PASSWORD etc. used by requisition)
def _load_env(path: str = ".env") -> None:
    env_file = Path(path)
    if not env_file.is_absolute():
        for base in (PROJECT_DIR, SCRIPT_DIR):
            candidate = base / env_file
            if candidate.exists():
                env_file = candidate
                break
    if not env_file.exists():
        return
    for raw in env_file.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, _, value = line.partition("=")
        key = key.strip()
        value = value.strip().strip("\"'")
        if key and key not in os.environ:
            os.environ[key] = value

_load_env()

import requisition as R
from playwright.sync_api import sync_playwright


# ===========================================================================
# COLUMN MAPS (legacy two-tab Excel)
# ===========================================================================

PR_COLUMNS = {
    "description": 1,
    "transaction_date": 2,
    "sub_type": 3,
    "category": 4,
    "site": 5,
    "account_site": 6,
    "party_code": 7,
    "party_desc": 8,
    "party_address": 9,
    "remarks": 10,
    "comments": 11,
    "status": 12,
}

ITEM_COLUMNS = {
    "pr_row": 1,
    "item_code": 2,
    "item_desc": 3,
    "ac_code": 4,
    "req_date": 5,
    "uom": 6,
    "pack_size": 7,
    "pack_qty": 8,
    "qty": 9,
    "base_uom": 10,
    "base_qty": 11,
    "rate": 12,
    "amount": 13,
    "memo": 14,
    "weight_uom": 15,
    "weight_qty": 16,
    "po_number": 17,
    "cost_center": 18,
    "remarks": 19,
    "loc_type": 20,
    "loc_id": 21,
}


# ===========================================================================
# FORM / DEPT / ITEM HELPERS (from form_pr.py)
# ===========================================================================

def _log_form(record: dict) -> None:
    try:
        LOG_DIR.mkdir(parents=True, exist_ok=True)
        with open(FORM_LOG, "a", encoding="utf-8") as f:
            f.write(json.dumps(record, ensure_ascii=False) + "\n")
    except Exception:
        pass


def _col(logical: str) -> str:
    """Map logical field name to sheet header via FORM_COLUMNS env."""
    cfg = os.getenv("FORM_COLUMNS", "indentor_name,item_name,item_quantity,department_name")
    parts = [p.strip() for p in cfg.split(",") if p.strip()]
    defaults = ["indentor_name", "item_name", "item_quantity", "department_name"]
    mapping = dict(zip(defaults, parts))
    return mapping.get(logical, logical)


_DEPT_SITE = {
    "dewas": "ADROIT DEWAS-Adroit Industries India Ltd",
    "driveshaft": "ADROIT DRIVESHAFT-Adroit Driveshafts Pvt Ltd",
    "drive": "ADROIT DRIVESHAFT-Adroit Driveshafts Pvt Ltd",
    "indore": "ADROIT INDORE-Adroit Industries India Ltd",
}


def _dept_to_site(dept: str) -> str:
    d = (dept or "").strip().lower()
    for k, v in _DEPT_SITE.items():
        if k in d:
            return v
    return os.getenv("FORM_DEFAULT_SITE", "ADROIT DEWAS-Adroit Industries India Ltd")


def _dept_to_account_site(dept: str) -> str:
    d = (dept or "").strip().lower()
    if "drive" in d:
        return "ADROIT DRIVESHAFT"
    if "indore" in d:
        return "ADROIT INDORE"
    return os.getenv("FORM_DEFAULT_ACCOUNT_SITE", "ADROIT DEWAS")


def _load_item_master() -> dict:
    p = Path(os.getenv("ITEM_MASTER_FILE", str(PROJECT_DIR / "data" / "item_master.json")))
    if p.is_file():
        try:
            return json.loads(p.read_text(encoding="utf-8"))
        except Exception:
            pass
    return {}


def _slug(value: str) -> str:
    s = (value or "").strip().upper()
    s = re.sub(r"[^A-Z0-9]+", "_", s)
    s = s.strip("_")
    return s[:40] or "ITEM"


def derive_item(item_name, item_qty="", department="", indentor="") -> dict:
    """Derive item_code / ac_code / uom from item name (optional item_master.json)."""
    master = _load_item_master()
    key = (item_name or "").strip().lower()
    base = master.get(key, {})
    return {
        "item_code": base.get("item_code") or _slug(item_name),
        "item_desc": base.get("item_desc") or item_name,
        "ac_code": base.get("ac_code") or "",
        "uom": base.get("uom") or os.getenv("FORM_DEFAULT_UOM", "NOS"),
        "qty": item_qty,
    }


def load_row(request_id: str) -> dict | None:
    """Read the sheet row whose RequestID column equals request_id."""
    from google.oauth2.service_account import Credentials
    import gspread

    sheet_id = os.getenv("GOOGLE_SHEET_ID")
    creds_file = os.getenv("GOOGLE_CREDS_FILE")
    if not sheet_id or not creds_file:
        raise RuntimeError("GOOGLE_SHEET_ID and GOOGLE_CREDS_FILE must be set")
    creds = Credentials.from_service_account_file(
        creds_file,
        scopes=["https://www.googleapis.com/auth/spreadsheets"],
    )
    client = gspread.authorize(creds)
    sh = client.open_by_key(sheet_id)
    ws = (
        sh.worksheet(os.getenv("FORM_SHEET_TAB", "Requests"))
        if os.getenv("FORM_SHEET_TAB")
        else sh.sheet1
    )
    records = ws.get_all_records()
    req_col = os.getenv("FORM_REQUEST_COLUMN", "RequestID")
    for r in records:
        if str(r.get(req_col, "")).strip() == str(request_id).strip():
            return r
    return None


def map_row_to_pr(row: dict) -> dict:
    """Map a form/Requests row to PR data dict."""
    indentor = (row.get(_col("indentor_name")) or "").strip()
    item_name = (row.get(_col("item_name")) or "").strip()
    qty = (row.get(_col("item_quantity")) or "").strip()
    dept = (row.get(_col("department_name")) or "").strip()

    item = derive_item(item_name, qty, dept, indentor)
    today = time.strftime("%d/%m/%Y")
    description = f"Indentor: {indentor} | Dept: {dept} | Item: {item_name}"

    return {
        "description": description,
        "transaction_date": today,
        "sub_type": os.getenv("FORM_DEFAULT_SUBTYPE", "Direct"),
        "category": os.getenv("FORM_DEFAULT_CATEGORY", "2627-PR-2627-PR"),
        "site": _dept_to_site(dept),
        "account_site": _dept_to_account_site(dept),
        "party_code": "",
        "party_desc": "",
        "party_address": "",
        "remarks": f"Created from web form. Indentor: {indentor}, Dept: {dept}",
        "comments": "",
        "items": [
            {
                "item_code": item["item_code"],
                "item_desc": item.get("item_desc", item_name),
                "ac_code": item.get("ac_code", ""),
                "req_date": today,
                "uom": item.get("uom", os.getenv("FORM_DEFAULT_UOM", "NOS")),
                "qty": qty,
                "base_uom": item.get("uom", os.getenv("FORM_DEFAULT_UOM", "NOS")),
                "base_qty": qty,
            }
        ],
    }


# ===========================================================================
# FLAT / EXCEL / GOOGLE SOURCE LOADING (from pipeline.py)
# ===========================================================================

def _find_flat_column_indices(headers):
    mapping = {}
    for idx, h in enumerate(headers or []):
        if h is None:
            continue
        hl = str(h).strip().lower()
        if not hl:
            continue
        if "item" in hl and "code" in hl:
            mapping.setdefault("item_code", idx)
            continue
        is_desc = False
        if "item" in hl:
            if any(x in hl for x in ("desc", "descrip", "desic", "descr", "name")):
                is_desc = True
            elif not any(x in hl for x in ("code", "date", "quant", "qty")):
                is_desc = True
        if is_desc:
            mapping.setdefault("item_desc", idx)
            continue
        if "quant" in hl or hl == "qty" or "qty" in hl:
            mapping.setdefault("qty", idx)
            continue
        if "date" in hl:
            mapping.setdefault("date", idx)
            continue
    return mapping


def _format_date_value(value):
    import datetime as _dt
    if value is None or (isinstance(value, str) and not value.strip()):
        return time.strftime("%d/%m/%Y")
    if isinstance(value, (_dt.datetime, _dt.date)):
        try:
            return value.strftime("%d/%m/%Y")
        except Exception:
            return str(value)
    s = str(value).strip()
    if not s:
        return time.strftime("%d/%m/%Y")
    for fmt in (
        "%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%m/%d/%Y",
        "%d.%m.%Y", "%d/%m/%y", "%Y/%m/%d",
    ):
        try:
            return _dt.datetime.strptime(s, fmt).strftime("%d/%m/%Y")
        except Exception:
            continue
    try:
        serial = float(s)
        if 30000 < serial < 60000:
            base = _dt.datetime(1899, 12, 30)
            return (base + _dt.timedelta(days=int(serial))).strftime("%d/%m/%Y")
    except Exception:
        pass
    return s


def _flat_row_to_pr(item_code, item_desc, date_str, qty_str, row_idx):
    item_code = str(item_code or "").strip()
    item_desc = str(item_desc or "").strip()
    if not item_code and item_desc:
        item_code = _slug(item_desc)
    if not item_desc and item_code:
        item_desc = item_code
    date_str = _format_date_value(date_str)
    qty_str = str(qty_str or "").strip() or "1"

    ac_code = ""
    uom = os.getenv("FORM_DEFAULT_UOM", "NOS")
    try:
        master = _load_item_master()
        for key in (item_desc.lower().strip(), item_code.lower().strip()):
            if key and key in master:
                base = master[key]
                ac_code = base.get("ac_code", ac_code)
                uom = base.get("uom", uom)
                break
    except Exception:
        pass

    return {
        "description": item_desc or f"PR for {item_code}",
        "transaction_date": date_str,
        "sub_type": os.getenv("FORM_DEFAULT_SUBTYPE", "Direct"),
        "category": os.getenv("FORM_DEFAULT_CATEGORY", "2627-PR-2627-PR"),
        "site": os.getenv("FORM_DEFAULT_SITE", "ADROIT DEWAS-Adroit Industries India Ltd"),
        "account_site": os.getenv("FORM_DEFAULT_ACCOUNT_SITE", "ADROIT DEWAS"),
        "party_code": "",
        "party_desc": "",
        "party_address": "",
        "remarks": f"Auto from flat sheet row {row_idx}: {item_desc}",
        "comments": "",
        "row": row_idx,
        "items": [
            {
                "item_code": item_code,
                "item_desc": item_desc,
                "ac_code": ac_code,
                "req_date": date_str,
                "uom": uom,
                "pack_size": "",
                "pack_qty": "",
                "qty": qty_str,
                "base_uom": uom,
                "base_qty": qty_str,
                "rate": "",
                "amount": "",
                "memo": "",
                "weight_uom": "",
                "weight_qty": "",
                "po_number": "",
                "cost_center": "",
                "remarks": "",
                "loc_type": "",
                "loc_id": "",
            }
        ],
    }


def _try_load_requests_excel(wb):
    sheet_name = None
    for cand in ("Requests", "requests", "REQUESTS"):
        if cand in wb.sheetnames:
            sheet_name = cand
            break
    if not sheet_name:
        for n in wb.sheetnames:
            if n.strip().lower() == "requests":
                sheet_name = n
                break
    if not sheet_name:
        return None
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        return None
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    hmap = {h.lower().strip(): idx for idx, h in enumerate(headers) if h}

    def _find_idx(*candidates):
        for cand in candidates:
            cand_l = cand.lower().strip()
            if cand_l in hmap:
                return hmap[cand_l]
            for hk, idx in hmap.items():
                if cand_l in hk:
                    return idx
        return None

    idx_item_name = _find_idx("item name", "item_name", "item")
    idx_item_code = _find_idx("item code", "item_code", "itemcode")
    idx_qty = _find_idx("item quantity", "item_quantity", "quantity", "qty", "required_quantity")
    idx_indentor = _find_idx("indenter name", "indentor name", "indenter", "indentor", "requester")
    idx_dept = _find_idx("department name", "department", "indenter dept", "indentor dept", "dept")
    idx_date = _find_idx("timestamp", "date", "requested at", "requested_at", "transaction_date", "req date")
    idx_for = _find_idx("item for", "description", "item for")

    if idx_item_name is None and idx_item_code is None:
        return None

    prs = []
    for r_idx, row in enumerate(rows[1:], start=2):
        if row is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
            continue

        def _val(idx):
            if idx is None or idx >= len(row):
                return ""
            v = row[idx]
            return "" if v is None else str(v).strip()

        item_name = _val(idx_item_name)
        item_code = _val(idx_item_code)
        qty = _val(idx_qty) or "1"
        indentor = _val(idx_indentor)
        dept = _val(idx_dept)
        date_val = _val(idx_date)
        item_for = _val(idx_for)

        if not item_name and not item_code:
            continue
        if not item_code and item_name:
            item_code = _slug(item_name)
        if not item_name and item_code:
            item_name = item_code
        try:
            qty_int = str(int(float(qty))) if qty else "1"
        except Exception:
            qty_int = qty or "1"

        pr = _flat_row_to_pr(item_code, item_name, date_val, qty_int, r_idx)
        if indentor or dept or item_for:
            pr["description"] = (
                f"Indentor: {indentor} | Dept: {dept} | Item: {item_name}"
                + (f" | For: {item_for}" if item_for else "")
            )
            pr["remarks"] = f"From Requests row {r_idx}: {indentor}/{dept} - {item_for or item_name}"
            pr["site"] = _dept_to_site(dept)
            pr["account_site"] = _dept_to_account_site(dept)
        if pr["items"]:
            pr["items"][0]["qty"] = qty_int
            pr["items"][0]["base_qty"] = qty_int
        prs.append(pr)
    if prs:
        print(f"Detected Requests sheet format in '{sheet_name}': {len(prs)} rows")
        return prs
    return None


def _try_load_flat_excel(wb):
    skip_names = {"fields reference", "field reference", "reference", "fields"}
    for name in wb.sheetnames:
        if name.strip().lower() in skip_names:
            continue
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows or len(rows) < 2:
            continue
        headers = rows[0]
        mapping = _find_flat_column_indices(headers)
        if ("item_code" not in mapping and "item_desc" not in mapping) or "date" not in mapping:
            continue
        prs = []
        for idx, row in enumerate(rows[1:], start=2):
            if row is None or all(v is None or (isinstance(v, str) and not v.strip()) for v in row):
                continue

            def _cell(key):
                if key not in mapping:
                    return ""
                col = mapping[key]
                return row[col] if col < len(row) else ""

            item_code = _cell("item_code")
            item_desc = _cell("item_desc")
            date_val = _cell("date")
            qty_val = _cell("qty")
            if not str(item_code or "").strip() and not str(item_desc or "").strip():
                continue
            prs.append(_flat_row_to_pr(item_code, item_desc, date_val, qty_val, idx))
        if prs:
            print(f"Detected flat Excel format in sheet '{name}': {len(prs)} rows")
            return prs
    return None


def _try_load_flat_google(worksheets):
    for title, ws in worksheets.items():
        try:
            values = ws.get_all_values()
        except Exception:
            continue
        if not values or len(values) < 2:
            continue
        headers = values[0]
        mapping = _find_flat_column_indices(headers)
        if ("item_code" not in mapping and "item_desc" not in mapping) or "date" not in mapping:
            continue
        prs = []
        for idx, row_vals in enumerate(values[1:], start=2):
            if not row_vals or all(not str(v or "").strip() for v in row_vals):
                continue

            def _gcell(key):
                if key not in mapping:
                    return ""
                col = mapping[key]
                return row_vals[col] if col < len(row_vals) else ""

            item_code = _gcell("item_code")
            item_desc = _gcell("item_desc")
            date_raw = _gcell("date")
            qty_raw = _gcell("qty")
            if not str(item_code or "").strip() and not str(item_desc or "").strip():
                continue
            prs.append(_flat_row_to_pr(item_code, item_desc, date_raw, qty_raw, idx))
        if prs:
            print(f"Detected flat Google Sheet format in tab '{title}': {len(prs)} rows")
            return prs
    return None


def _is_google_sheet_source(source):
    if not source:
        return False
    src_str = str(source).strip()
    if src_str.startswith(("http://", "https://", "gsheet://")):
        return True
    if src_str.endswith((".xlsx", ".xls", ".csv")):
        return False
    path = Path(src_str)
    if not path.is_absolute():
        path = DATA_DIR / path
    if path.is_file():
        return False
    if len(src_str) >= 20 or src_str.lower() in ("gsheet", "google"):
        return True
    return False


def _load_from_google_sheet(source=None, creds_file=None):
    from google.oauth2.service_account import Credentials
    import gspread

    sheet_id_or_url = source
    if not sheet_id_or_url or str(sheet_id_or_url).lower() in ("gsheet", "google"):
        sheet_id_or_url = os.getenv("GOOGLE_SHEET_ID")
    if not sheet_id_or_url:
        print("ERROR: No Google Sheet ID/URL and GOOGLE_SHEET_ID not set.")
        return []

    sheet_id_or_url = str(sheet_id_or_url).strip()
    if sheet_id_or_url.startswith("gsheet://"):
        sheet_id_or_url = sheet_id_or_url.replace("gsheet://", "")

    if not creds_file:
        creds_file = os.getenv(
            "GOOGLE_CREDS_FILE",
            str(PROJECT_DIR / "credentials" / "service_account.json"),
        )
    creds_path = Path(creds_file)
    if not creds_path.is_absolute():
        creds_path = PROJECT_DIR / creds_path
    if not creds_path.is_file():
        print(f"ERROR: Service account file {creds_path} not found.")
        return []

    creds = Credentials.from_service_account_file(
        str(creds_path),
        scopes=["https://www.googleapis.com/auth/spreadsheets"],
    )
    client = gspread.authorize(creds)
    if sheet_id_or_url.startswith("http://") or sheet_id_or_url.startswith("https://"):
        sh = client.open_by_url(sheet_id_or_url)
    else:
        sh = client.open_by_key(sheet_id_or_url)

    worksheets = {ws.title: ws for ws in sh.worksheets()}
    item_sheet_name = (
        "PR Items" if "PR Items" in worksheets
        else ("PR Input" if "PR Input" in worksheets else None)
    )

    # Legacy two-tab
    if "PR" in worksheets and item_sheet_name:
        wsi = worksheets[item_sheet_name]
        item_data = wsi.get_all_values()
        item_rows = []
        for row_vals in item_data[1:]:
            if not row_vals or not row_vals[0]:
                continue
            try:
                pr_row_val = int(row_vals[0])
            except ValueError:
                continue
            items = {}
            for key, col in ITEM_COLUMNS.items():
                val = row_vals[col - 1] if (col - 1) < len(row_vals) else ""
                items[key] = val
            item_rows.append((pr_row_val, items))

        ws_pr = worksheets["PR"]
        pr_data = ws_pr.get_all_values()
        prs = []
        for idx, row_vals in enumerate(pr_data[1:], start=2):
            if not row_vals or (not row_vals[0] and (len(row_vals) <= 2 or not row_vals[2])):
                continue
            data = {}
            for key, col in PR_COLUMNS.items():
                val = row_vals[col - 1] if (col - 1) < len(row_vals) else ""
                data[key] = val
            data["row"] = idx
            data["items"] = [it for r, it in item_rows if r == idx]
            prs.append(data)
        if prs and any(p.get("items") for p in prs):
            return prs
        if prs:
            print(f"Google PR sheets produced {len(prs)} rows but empty items — trying flat")

    flat = _try_load_flat_google(worksheets)
    if flat:
        return flat

    form_tab = os.getenv("FORM_SHEET_TAB", "Requests")
    if form_tab in worksheets or worksheets:
        ws_form = worksheets.get(form_tab) or list(worksheets.values())[0]
        records = ws_form.get_all_records()
        prs = []
        for idx, rec in enumerate(records, start=2):
            pr_data = map_row_to_pr(rec)
            pr_data["row"] = idx
            prs.append(pr_data)
        return prs

    print("WARNING: Google Sheet needs PR+PR Items tabs, flat columns, or form tab.")
    return []


def load_pr_source(source=None):
    """Load list of PR dicts from Excel or Google Sheet."""
    if source and _is_google_sheet_source(source):
        print(f"Loading PR data from Google Sheet: {source}")
        return _load_from_google_sheet(source)

    import openpyxl

    path = Path(source or "PR_Input.xlsx")
    if not path.is_absolute():
        path = DATA_DIR / path

    if not path.is_file():
        if DATA_DIR.is_dir() and path.parent == DATA_DIR:
            wanted = path.name.lower()
            for cand in DATA_DIR.iterdir():
                if cand.name.lower() == wanted:
                    path = cand
                    break
        if not path.is_file():
            gsheet_id = os.getenv("GOOGLE_SHEET_ID")
            if gsheet_id:
                print(f"Local file {path} not found. Falling back to GOOGLE_SHEET_ID...")
                return _load_from_google_sheet(gsheet_id)
            print(f"WARNING: Source file {path} does not exist.")
            return []

    wb = openpyxl.load_workbook(path, data_only=True)

    # Legacy two-tab
    item_sheet = (
        "PR Items" if "PR Items" in wb.sheetnames
        else ("PR Input" if "PR Input" in wb.sheetnames else None)
    )
    if "PR" in wb.sheetnames and item_sheet is not None:
        wsi = wb[item_sheet]
        item_rows = []
        for row in wsi.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            try:
                pr_row_idx = int(row[0])
            except Exception:
                continue
            items = {
                key: (row[col - 1] if (col - 1) < len(row) else "")
                for key, col in ITEM_COLUMNS.items()
            }
            item_rows.append((pr_row_idx, items))

        prs = []
        ws = wb["PR"]
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if row[0] is None and (len(row) <= 2 or row[2] is None):
                continue
            data = {
                key: (row[col - 1] if (col - 1) < len(row) else "")
                for key, col in PR_COLUMNS.items()
            }
            data["row"] = idx
            data["items"] = [it for r, it in item_rows if r == idx]
            prs.append(data)
        if prs and any(p.get("items") for p in prs):
            return prs
        if prs:
            print(f"PR sheets produced {len(prs)} rows but empty items — trying flat")

    req_prs = _try_load_requests_excel(wb)
    if req_prs:
        return req_prs

    flat = _try_load_flat_excel(wb)
    if flat:
        return flat

    print(
        f"WARNING: {path.name} needs 'PR'+'PR Items'/'PR Input', "
        "or a flat sheet with item Code/Description + Date columns"
    )
    return []


# ===========================================================================
# ONE-SHOT PR (item-code + qty)
# ===========================================================================

def build_one_shot_pr(item_code: str, qty, **overrides) -> dict:
    today = datetime.now().strftime("%d/%m/%Y")
    qty_str = str(qty)
    return {
        "description": overrides.get("description") or f"PR for {item_code}",
        "transaction_date": today,
        "sub_type": overrides.get("sub_type") or os.getenv("FORM_DEFAULT_SUBTYPE", "Direct"),
        "category": overrides.get("category") or os.getenv("FORM_DEFAULT_CATEGORY", "2627-PR-2627-PR"),
        "site": overrides.get("site") or os.getenv(
            "FORM_DEFAULT_SITE", "ADROIT DEWAS-Adroit Industries India Ltd"
        ),
        "account_site": overrides.get("account_site") or os.getenv(
            "FORM_DEFAULT_ACCOUNT_SITE", "ADROIT DEWAS"
        ),
        "party_code": "",
        "party_desc": "",
        "party_address": "",
        "remarks": overrides.get("remarks") or f"Auto PR for {item_code} qty {qty_str}",
        "comments": "",
        "items": [
            {
                "item_code": item_code,
                "item_desc": item_code,
                "ac_code": "",
                "req_date": today,
                "uom": os.getenv("FORM_DEFAULT_UOM", "NOS"),
                "pack_size": "",
                "pack_qty": "",
                "qty": qty_str,
                "base_uom": os.getenv("FORM_DEFAULT_UOM", "NOS"),
                "base_qty": qty_str,
                "rate": "",
                "amount": "",
            }
        ],
    }


def create_and_approve_pr(page, pr_data) -> tuple[str | None, bool]:
    """Navigate to requisition form, fill, save draft, approve. Returns (req_no, approved)."""
    if not R.navigate_to_requisition_form(page):
        print("FAIL: could not reach requisition form")
        return None, False

    req_no = R.fill_and_save_draft(page, pr_data)
    if not req_no:
        print("FAIL: could not save requisition draft")
        return None, False
    print(f"  Requisition saved: {req_no}")

    approved = False
    if hasattr(R, "approve_if_not_approved"):
        approved = R.approve_if_not_approved(page, req_no)
    if not approved and hasattr(R, "submit_draft"):
        print("  Trying submit_draft...")
        approved = bool(R.submit_draft(page, req_no))
    if approved:
        print(f"  Requisition {req_no} approved")
    else:
        print(f"  WARNING: could not confirm approval for {req_no}")
    return req_no, approved


# ===========================================================================
# PIPELINE (batch)
# ===========================================================================

def run_pipeline(source=None, row=None) -> list:
    prs = load_pr_source(source or "PR_Input.xlsx")
    if row is not None:
        prs = [p for p in prs if p.get("row") == row]
    if not prs:
        print("No PR rows to process")
        return []

    results = []
    os.makedirs(getattr(R, "SCREENSHOT_DIR", SCREENSHOT_DIR), exist_ok=True)

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False, args=["--start-maximized"])
        context = browser.new_context(no_viewport=True)
        page = context.new_page()
        try:
            if not R.do_login(page):
                return [{"row": None, "result": "login_failed"}]
            for data in prs:
                print(f"\n===== Processing PR sheet row {data.get('row')} =====")
                try:
                    if page.is_closed():
                        print("  Previous flow closed the tab; opening a fresh page...")
                        page = context.new_page()
                        if not R.do_login(page):
                            results.append({"row": data.get("row"), "result": "relogin_failed"})
                            continue
                    req_no, approved = create_and_approve_pr(page, data)
                    if not req_no:
                        results.append({"row": data.get("row"), "result": "save_failed"})
                        continue
                    results.append({
                        "row": data.get("row"),
                        "requisition_no": req_no,
                        "submitted": approved,
                    })
                except Exception as exc:
                    print(f"  Row {data.get('row')} errored: {exc}")
                    results.append({
                        "row": data.get("row"),
                        "result": "error",
                        "error": str(exc)[:200],
                    })
        except KeyboardInterrupt:
            print("Stopped by user.")
        finally:
            R.safe_logout(page)
            try:
                browser.close()
            except Exception:
                pass
    return results


# ===========================================================================
# FORM REQUEST (single token)
# ===========================================================================

def run_form_pr(request_id: str) -> dict:
    _log_form({"event": "form_pr_start", "request_id": request_id})
    try:
        row = load_row(request_id)
    except Exception as exc:
        result = {"request_id": request_id, "result": "sheet_error", "error": str(exc)}
        _log_form(result)
        return result
    if not row:
        result = {"request_id": request_id, "result": "not_found"}
        _log_form(result)
        return result

    data = map_row_to_pr(row)
    os.makedirs(getattr(R, "SCREENSHOT_DIR", SCREENSHOT_DIR), exist_ok=True)

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False, args=["--start-maximized"])
        context = browser.new_context(no_viewport=True)
        page = context.new_page()
        try:
            if not R.do_login(page):
                result = {"request_id": request_id, "result": "login_failed"}
                _log_form(result)
                return result
            req_no, approved = create_and_approve_pr(page, data)
            if not req_no:
                result = {"request_id": request_id, "result": "save_failed"}
                _log_form(result)
                return result
            result = {
                "request_id": request_id,
                "requisition_no": req_no,
                "submitted": bool(approved),
                "result": "ok",
            }
            _log_form(result)
            return result
        except Exception as exc:
            result = {"request_id": request_id, "result": "error", "error": str(exc)}
            _log_form(result)
            return result
        finally:
            R.safe_logout(page)
            try:
                browser.close()
            except Exception:
                pass


# ===========================================================================
# JSON PR MODE (worker contract: create_pr)
# ===========================================================================

def _norm_qty(value) -> str:
    """Normalize a quantity value to its canonical string form ('50.0' -> '50')."""
    v = str(value or 1).strip()
    if not v:
        return "1"
    try:
        f = float(v)
        return str(int(f)) if f.is_integer() else str(f)
    except ValueError:
        return v


def build_pr_from_json(payload: dict) -> dict:
    """Map the create_pr JSON payload to a PR data dict.

    Contract fields (aliases accepted):
      itemCode / item_code             (required)
      itemQuantity / item_quantity / quantity / qty
      vendorCode / vendor_code
      site
      date / transactionDate / transaction_date   (dd/MM/yyyy)
      uom
      requestId / request_id
      itemDesc / itemDescription / description
    """
    item_code = str(payload.get("itemCode") or payload.get("item_code") or "").strip()
    if not item_code:
        raise ValueError("itemCode is required")
    qty = _norm_qty(
        payload.get("itemQuantity")
        or payload.get("item_quantity")
        or payload.get("quantity")
        or payload.get("qty")
        or 1
    )
    vendor = str(payload.get("vendorCode") or payload.get("vendor_code") or "").strip()
    request_id = str(payload.get("requestId") or payload.get("request_id") or "").strip()
    site = str(payload.get("site") or "").strip()
    uom = str(payload.get("uom") or "").strip()
    item_desc = str(
        payload.get("itemDesc")
        or payload.get("itemDescription")
        or payload.get("description")
        or ""
    ).strip()
    date_raw = payload.get("date") or payload.get("transactionDate") or payload.get("transaction_date")

    data = build_one_shot_pr(item_code, qty)
    if item_desc:
        data["description"] = item_desc
        data["items"][0]["item_desc"] = item_desc
    if date_raw:
        data["transaction_date"] = _format_date_value(date_raw)
        data["items"][0]["req_date"] = data["transaction_date"]
    if vendor:
        data["party_code"] = vendor
        data["party_desc"] = vendor
    if uom:
        data["items"][0]["uom"] = uom
        data["items"][0]["base_uom"] = uom
    if site:
        data["site"] = site
        s_low = (site or "").lower()
        if "driveshaft" in s_low:
            data["account_site"] = "ADROIT DRIVESHAFT"
        elif "indore" in s_low:
            data["account_site"] = "ADROIT INDORE"
        elif os.getenv("FORM_DEFAULT_ACCOUNT_SITE"):
            data["account_site"] = os.getenv("FORM_DEFAULT_ACCOUNT_SITE")
    data["remarks"] = (data.get("remarks") or "") + (
        f" | RequestID: {request_id}" if request_id else ""
    )
    data["request_id"] = request_id
    return data


def _load_json_input(path: str) -> dict:
    p = Path(path)
    if not p.is_absolute():
        for base in (PROJECT_DIR, SCRIPT_DIR, DATA_DIR):
            cand = base / p
            if cand.is_file():
                p = cand
                break
    if not p.is_file():
        raise SystemExit(f"JSON input file not found: {path}")
    return json.loads(p.read_text(encoding="utf-8"))


def run_json_pr(payload: dict) -> dict:
    """Run a single PR from the create_pr JSON contract. Result JSON = last stdout line."""
    started = time.time()
    item_code = str(payload.get("itemCode") or payload.get("item_code") or "").strip()
    vendor = str(payload.get("vendorCode") or payload.get("vendor_code") or "").strip()
    request_id = str(payload.get("requestId") or payload.get("request_id") or "").strip()
    qty = _norm_qty(
        payload.get("itemQuantity")
        or payload.get("item_quantity")
        or payload.get("quantity")
        or payload.get("qty")
        or 1
    )

    result = {
        "ok": False,
        "action": "create_pr",
        "prNumber": None,
        "status": None,
        "itemCode": item_code,
        "qty": qty,
        "requestId": request_id,
        "vendorCode": vendor,
        "transactionDate": _format_date_value(
            payload.get("date")
            or payload.get("transactionDate")
            or payload.get("transaction_date")
        ),
        "durationMs": None,
        "error": None,
    }

    def _fail(msg):
        result["error"] = str(msg)[:500]
        result["durationMs"] = int((time.time() - started) * 1000)
        result["status"] = "failed"
        return result

    try:
        data = build_pr_from_json(payload)
    except Exception as exc:
        return _fail(exc)

    os.makedirs(getattr(R, "SCREENSHOT_DIR", SCREENSHOT_DIR), exist_ok=True)
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False, args=["--start-maximized"])
        context = browser.new_context(no_viewport=True)
        page = context.new_page()
        try:
            if not R.do_login(page):
                return _fail("login failed")
            req_no, approved = create_and_approve_pr(page, data)
            if not req_no:
                return _fail("save failed / no requisition number returned")
            result["ok"] = True
            result["prNumber"] = req_no
            result["status"] = "approved" if approved else "saved"
            result["durationMs"] = int((time.time() - started) * 1000)
            return result
        except Exception as exc:
            return _fail(exc)
        finally:
            try:
                R.safe_logout(page)
            except Exception:
                pass
            try:
                browser.close()
            except Exception:
                pass


# ===========================================================================
# TCS WORKER (Sheets queue poller — PR only)
# ===========================================================================

TCS_REQUISITION_NO_HDR = "TCS Requisition No"
TCS_STATUS_HDR = "TCS Status"
TCS_UPDATED_HDR = "TCS Updated At"
TCS_ERROR_HDR = "TCS Error"
TCS_SCREENSHOTS_HDR = "TCS Screenshots"

PENDING_TCS_STATUSES = {"", "pending", "queued", "retry", "null", "none"}


def _worker_log(msg):
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    line = f"[tcs-worker {ts}] {msg}"
    print(line, flush=True)
    try:
        log_path = LOG_DIR / "tcs_worker.log"
        log_path.parent.mkdir(parents=True, exist_ok=True)
        with open(log_path, "a", encoding="utf-8") as f:
            f.write(line + "\n")
    except Exception:
        pass


def _sheet_client():
    try:
        import gspread
        from google.oauth2.service_account import Credentials
    except ImportError as e:
        _worker_log(f"Missing deps: {e} — pip install gspread google-auth")
        return None, None

    creds_file = os.getenv(
        "GOOGLE_CREDS_FILE",
        str(PROJECT_DIR / "credentials" / "service_account.json"),
    )
    sheet_id = os.getenv("GOOGLE_SHEET_ID", "").strip()
    if not sheet_id:
        _worker_log("ERROR: GOOGLE_SHEET_ID empty")
        return None, None
    creds_path = Path(creds_file)
    if not creds_path.is_absolute():
        creds_path = PROJECT_DIR / creds_path
    if not creds_path.is_file():
        _worker_log(f"ERROR: service account {creds_path} missing")
        return None, None

    creds = Credentials.from_service_account_file(
        str(creds_path),
        scopes=["https://www.googleapis.com/auth/spreadsheets"],
    )
    client = gspread.authorize(creds)
    if sheet_id.startswith("http"):
        sh = client.open_by_url(sheet_id)
    else:
        sh = client.open_by_key(sheet_id)
    return client, sh


def _header_map(ws):
    headers = ws.row_values(1)
    m = {}
    for idx, h in enumerate(headers):
        k = str(h or "").strip().lower()
        if k:
            m[k] = idx
    return m, headers


def _ensure_tcs_columns(ws):
    m, headers = _header_map(ws)
    needed = [
        TCS_REQUISITION_NO_HDR,
        TCS_STATUS_HDR,
        TCS_UPDATED_HDR,
        TCS_ERROR_HDR,
        TCS_SCREENSHOTS_HDR,
    ]
    missing = [n for n in needed if n.lower() not in m]
    if missing:
        _worker_log(f"Adding missing columns to {ws.title}: {missing}")
        next_col = len(headers) + 1
        for i, col_name in enumerate(missing):
            ws.update_cell(1, next_col + i, col_name)
            time.sleep(0.3)
        m, headers = _header_map(ws)
    return m, headers


def _find_pending_tokens(sh):
    import gspread
    results = []
    for tab_name in [os.getenv("FORM_SHEET_TAB", "Requests"), "Approvals", "Requests"]:
        try:
            ws = sh.worksheet(tab_name)
        except gspread.WorksheetNotFound:
            continue
        _ensure_tcs_columns(ws)
        try:
            values = ws.get_all_values()
        except Exception as e:
            _worker_log(f"get_all_values {tab_name} failed: {e}")
            continue
        if len(values) < 2:
            continue
        hdr = [str(h).strip().lower() for h in values[0]]
        t_idx = next(
            (hdr.index(h) for h in hdr if h in ("token", "requestid", "request_id")),
            None,
        )
        s_idx = hdr.index("status") if "status" in hdr else None
        tcs_s_idx = next(
            (hdr.index(h) for h in hdr if h == TCS_STATUS_HDR.lower()),
            None,
        )
        if t_idx is None or s_idx is None:
            continue
        for row_idx, row in enumerate(values[1:], start=2):
            token = str(row[t_idx] or "").strip() if t_idx < len(row) else ""
            status = str(row[s_idx] or "").strip().lower() if s_idx < len(row) else ""
            tcs_status = (
                str(row[tcs_s_idx] or "").strip().lower()
                if tcs_s_idx is not None and tcs_s_idx < len(row)
                else ""
            )
            if not token or token.lower().startswith("requestid"):
                continue
            if status == "approved" and tcs_status in PENDING_TCS_STATUSES:
                if token not in [r["token"] for r in results]:
                    results.append({"token": token, "row_idx": row_idx, "tab": tab_name})
        if results and tab_name == "Approvals":
            break
    return results


def _update_tcs_status(sh, token, requisition_no="", tcs_status="", error="", screenshots=""):
    import gspread
    now_iso = datetime.now().isoformat()
    for tab_name in ["Requests", "Approvals"]:
        try:
            ws = sh.worksheet(tab_name)
        except gspread.WorksheetNotFound:
            continue
        m, headers = _header_map(ws)
        token_col = None
        for cand in ("token", "requestid", "request_id"):
            if cand in m:
                token_col = m[cand] + 1
                if tab_name == "Approvals" and cand == "token":
                    break
                if tab_name == "Requests" and cand in ("requestid", "request_id"):
                    break
        if not token_col:
            continue
        try:
            col_vals = ws.col_values(token_col)
        except Exception:
            continue
        row_idx = None
        for idx, val in enumerate(col_vals[1:], start=2):
            if str(val).strip().lower() == str(token).strip().lower():
                row_idx = idx
                break
        if not row_idx:
            continue
        m, headers = _ensure_tcs_columns(ws)

        def set_cell(hdr_name, value):
            key = hdr_name.lower()
            if key in m:
                col = m[key] + 1
                ws.update_cell(row_idx, col, value)
                time.sleep(0.2)

        if requisition_no is not None:
            set_cell(TCS_REQUISITION_NO_HDR, requisition_no)
        if tcs_status is not None:
            set_cell(TCS_STATUS_HDR, tcs_status)
        set_cell(TCS_UPDATED_HDR, now_iso)
        if error is not None:
            set_cell(TCS_ERROR_HDR, (error or "")[:500])
        if screenshots:
            set_cell(TCS_SCREENSHOTS_HDR, screenshots)


def run_worker_once(dry_run=False) -> list:
    client, sh = _sheet_client()
    if not sh:
        _worker_log("No sheet client — aborting")
        return []
    pending = _find_pending_tokens(sh)
    if not pending:
        _worker_log("No pending TCS jobs (approved but tcs_status empty/pending)")
        return []
    _worker_log(f"Found {len(pending)} pending token(s): {[p['token'] for p in pending]}")
    results = []
    if dry_run:
        for p in pending:
            _worker_log(f"[dry-run] would process {p['token']} from {p['tab']}")
            results.append({"token": p["token"], "dry_run": True})
        return results

    os.makedirs(getattr(R, "SCREENSHOT_DIR", SCREENSHOT_DIR), exist_ok=True)
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=False, args=["--start-maximized"])
        context = browser.new_context(no_viewport=True)
        page = context.new_page()
        try:
            if not R.do_login(page):
                _worker_log("TCS login failed — marking all pending as failed")
                for pend in pending:
                    _update_tcs_status(sh, pend["token"], tcs_status="failed", error="TCS login failed")
                    results.append({
                        "token": pend["token"],
                        "requisition_no": None,
                        "submitted": False,
                        "error": "login_failed",
                    })
                return results

            for pend in pending:
                token = pend["token"]
                _worker_log(f"Processing {token} ...")
                try:
                    _update_tcs_status(sh, token, tcs_status="processing", error="")
                    row = load_row(token)
                    if not row:
                        msg = f"Requests row not found for {token}"
                        _worker_log(msg)
                        _update_tcs_status(sh, token, tcs_status="failed", error=msg)
                        results.append({"token": token, "error": msg})
                        continue
                    data = map_row_to_pr(row)
                    if page.is_closed():
                        page = context.new_page()
                        if not R.do_login(page):
                            raise RuntimeError("relogin failed")
                    req_no, approved = create_and_approve_pr(page, data)
                    if not req_no:
                        raise RuntimeError("fill_and_save_draft returned no requisition_no")
                    _worker_log(f"Saved draft {req_no} for {token}")
                    final_status = "approved" if approved else "saved"
                    shots = (
                        "requisition_saved.png,req_after_approve.png"
                        if approved
                        else "requisition_saved.png"
                    )
                    _update_tcs_status(
                        sh, token,
                        requisition_no=req_no,
                        tcs_status=final_status,
                        screenshots=shots,
                    )
                    results.append({
                        "token": token,
                        "requisition_no": req_no,
                        "submitted": approved,
                    })
                    _log_form({
                        "token": token,
                        "requisition_no": req_no,
                        "submitted": approved,
                        "result": "ok",
                        "source": "tcs_worker",
                    })
                except Exception as e:
                    err = str(e)[:500]
                    _worker_log(f"Failed {token}: {err}")
                    traceback.print_exc()
                    _update_tcs_status(sh, token, tcs_status="failed", error=err)
                    results.append({"token": token, "error": err})
                    try:
                        if page.is_closed():
                            page = context.new_page()
                            R.do_login(page)
                    except Exception:
                        pass
        finally:
            try:
                R.safe_logout(page)
            except Exception:
                pass
            try:
                browser.close()
            except Exception:
                pass
    return results


# ===========================================================================
# CLI
# ===========================================================================

def main() -> int:
    parser = argparse.ArgumentParser(
        description="TCS iON Purchase Requisition (PR) bot — combined"
    )
    # One-shot
    parser.add_argument("--item-code", help="Item code for one-shot PR")
    parser.add_argument("--qty", type=int, help="Quantity for one-shot PR")
    # Pipeline
    parser.add_argument(
        "--source",
        default=None,
        help="Excel file, Google Sheet URL/ID, or 'gsheet'",
    )
    parser.add_argument("--row", type=int, default=None, help="Process only this sheet row")
    # Form
    parser.add_argument("--request-id", help="Form/Requests RequestID token")
    # Worker
    parser.add_argument("--worker", action="store_true", help="Run Sheets queue poller")
    parser.add_argument("--interval", type=int, default=0, help="Worker poll interval (0=once)")
    parser.add_argument("--once", action="store_true", help="Worker: run once")
    parser.add_argument("--dry-run", action="store_true", help="Worker: log only, no TCS portal")
    parser.add_argument(
        "--json",
        default=None,
        help="JSON input file (create_pr contract: action/itemCode/qty/vendorCode/requestId)",
    )

    args = parser.parse_args()

    # Worker mode
    if args.worker:
        if args.once:
            args.interval = 0
        if args.interval <= 0:
            res = run_worker_once(dry_run=args.dry_run)
            _worker_log(f"Run once done: {res}")
            return 0
        _worker_log(f"Starting poll loop every {args.interval}s (Ctrl+C to stop)")
        while True:
            try:
                res = run_worker_once(dry_run=args.dry_run)
                if res:
                    _worker_log(f"Batch done: {res}")
            except KeyboardInterrupt:
                _worker_log("Interrupted — exiting")
                break
            except Exception as e:
                _worker_log(f"Loop error: {e}")
                traceback.print_exc()
            time.sleep(args.interval)
        return 0

    # Form request-id mode
    if args.request_id:
        result = run_form_pr(args.request_id)
        print(f"\nForm PR result: {result}")
        return 0 if result.get("result") == "ok" else 1

    # JSON contract mode (worker: create_pr)
    if args.json:
        payload = _load_json_input(args.json)
        result = run_json_pr(payload)
        print(json.dumps(result, ensure_ascii=False))
        _log_form({
            "event": "json_pr",
            "ok": result.get("ok"),
            "request_id": result.get("requestId"),
            "pr_number": result.get("prNumber"),
            "status": result.get("status"),
            "error": result.get("error"),
        })
        return 0 if result.get("ok") else 1

    # One-shot item-code mode
    if args.item_code:
        if args.qty is None:
            print("ERROR: --qty is required with --item-code")
            return 1
        pr_data = build_one_shot_pr(args.item_code, args.qty)
        os.makedirs(getattr(R, "SCREENSHOT_DIR", SCREENSHOT_DIR), exist_ok=True)
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=False, args=["--start-maximized"])
            context = browser.new_context(no_viewport=True)
            page = context.new_page()
            try:
                print("\n[1/2] Logging in...")
                if not R.do_login(page):
                    print("FAIL: login failed")
                    return 1
                print(f"\n[2/2] Creating Requisition for {args.item_code} (qty {args.qty})...")
                req_no, approved = create_and_approve_pr(page, pr_data)
                print("\n" + "=" * 50)
                if req_no:
                    print(f"  Requisition: {req_no}")
                    print(f"  Approved: {approved}")
                    print(f"  Item: {args.item_code}, Qty: {args.qty}")
                    print("  RESULT: SUCCESS" if approved else "  RESULT: SAVED (approval uncertain)")
                else:
                    print("  RESULT: FAILED")
                print("=" * 50)
                return 0 if req_no else 1
            except KeyboardInterrupt:
                print("\nInterrupted by user.")
                return 130
            except Exception as e:
                print(f"\nERROR: {e}")
                traceback.print_exc()
                return 1
            finally:
                R.safe_logout(page)
                try:
                    browser.close()
                except Exception:
                    pass

    # Pipeline / source mode (default if --source given, else try PR_Input.xlsx)
    source = args.source or "PR_Input.xlsx"
    result = run_pipeline(source=source, row=args.row)
    print(f"\nPipeline result: {result}")
    return 0 if result and all(r.get("requisition_no") for r in result if "result" not in r or r.get("result") not in ("login_failed", "save_failed", "error")) else (0 if result else 1)


if __name__ == "__main__":
    sys.exit(main())
